from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import NamedTuple

import ebooklib
import fitz
import pytesseract
from bs4 import BeautifulSoup
from docx import Document
from ebooklib import epub
from PIL import Image, ImageOps

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".epub", ".png", ".jpg", ".jpeg", ".webp", ".html", ".htm", ".md", ".markdown", ".rtf"}
MAX_IMAGE_PIXELS = 40_000_000

class ExtractionResult(NamedTuple):
    text: str
    page_count: int
    format: str

def _clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")

def extract_document(data: bytes, filename: str, max_pdf_pages: int = 800, max_ocr_pages: int = 100) -> ExtractionResult:
    extension = Path(filename.replace("\\", "/")).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file format.")
    if not data:
        raise ValueError("The selected document is empty.")

    if extension == ".pdf":
        try:
            with fitz.open(stream=data, filetype="pdf") as pdf:
                if pdf.is_encrypted:
                    raise ValueError("Password-protected PDFs are not supported in this version.")
                if pdf.page_count > max_pdf_pages:
                    raise ValueError("PDF exceeds the " + str(max_pdf_pages) + "-page processing limit.")
                pages: list[str] = []
                ocr_count = 0
                for page in pdf:
                    page_text = page.get_text("text").strip()
                    if not page_text and ocr_count < max_ocr_pages:
                        try:
                            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                            image = Image.open(io.BytesIO(pix.tobytes("png")))
                            page_text = pytesseract.image_to_string(image).strip()
                            ocr_count += 1
                        except Exception:
                            page_text = ""
                    pages.append(page_text)
                return ExtractionResult(_clean_text("\n\n".join(pages)), pdf.page_count, "pdf")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("The PDF could not be opened or parsed.") from exc

    if extension == ".docx":
        try:
            document = Document(io.BytesIO(data))
            blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            text = _clean_text("\n\n".join(blocks))
            return ExtractionResult(text, max(1, text.count("\n\n") + 1) if text else 0, "docx")
        except Exception as exc:
            raise ValueError("The Word document could not be opened or parsed.") from exc

    if extension == ".epub":
        try:
            with tempfile.TemporaryDirectory(prefix="readaloud-") as temp_dir:
                path = Path(temp_dir) / "book.epub"
                path.write_bytes(data)
                book = epub.read_epub(str(path), options={"ignore_ncx": True})
                blocks: list[str] = []
                used: set[str] = set()
                for entry in getattr(book, "spine", []):
                    item_id = entry[0] if isinstance(entry, tuple) else entry
                    item = book.get_item_with_id(item_id)
                    if item is not None and item.get_type() == ebooklib.ITEM_DOCUMENT:
                        soup = BeautifulSoup(item.get_content(), "html.parser")
                        for tag in soup(["script", "style", "nav"]):
                            tag.decompose()
                        text = soup.get_text(" ", strip=True)
                        if text:
                            blocks.append(text)
                        used.add(item.get_id())
                for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                    if item.get_id() not in used:
                        soup = BeautifulSoup(item.get_content(), "html.parser")
                        for tag in soup(["script", "style", "nav"]):
                            tag.decompose()
                        text = soup.get_text(" ", strip=True)
                        if text:
                            blocks.append(text)
            joined = _clean_text("\n\n".join(blocks))
            return ExtractionResult(joined, max(1, len(blocks)) if joined else 0, "epub")
        except Exception as exc:
            raise ValueError("The EPUB book could not be opened or parsed.") from exc

    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with Image.open(io.BytesIO(data)) as source:
                if source.width * source.height > MAX_IMAGE_PIXELS:
                    raise ValueError("Image is too large to OCR safely.")
                image = ImageOps.exif_transpose(source).convert("RGB")
                text = pytesseract.image_to_string(image)
            return ExtractionResult(_clean_text(text), 1, "image")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Image OCR failed. Confirm Tesseract is installed and try a clearer image.") from exc

    text = _decode_text(data)
    if extension in {".html", ".htm"}:
        soup = BeautifulSoup(text, "html.parser")
        for tag in soup(["script", "style", "noscript", "nav"]):
            tag.decompose()
        text = soup.get_text("\n", strip=True)
    elif extension == ".rtf":
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?|[{}]", " ", text)
    result = _clean_text(text)
    return ExtractionResult(result, max(1, result.count("\n\n") + 1) if result else 0, extension.lstrip("."))
